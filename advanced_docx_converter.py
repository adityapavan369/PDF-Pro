""" 
Advanced DOCX to PDF Converter with Page Fitting Capabilities
This module provides enhanced conversion functionality to achieve single-page PDF output
when possible, while preserving formatting quality. This improved version is more robust
and includes tool availability checks, safer python-docx manipulations, better temp file
handling, and multiple fallbacks to increase success rates across environments.
"""
import subprocess
import os
import tempfile
import shutil
import logging
from pathlib import Path
from typing import Optional

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class AdvancedDocxConverter:
    """Advanced DOCX to PDF converter with multiple strategies for optimal output."""
    
    def __init__(self, timeout: int = 120):
        self.timeout = timeout
    
    # Helper utilities
    def _is_tool_available(self, name: str) -> bool:
        """Return True if `name` is found on PATH."""
        from shutil import which
        return which(name) is not None
    
    def _run_command(self, cmd, timeout: Optional[int] = None):
        """Run a subprocess command and return CompletedProcess. Raises on timeout."""
        try:
            to = timeout or self.timeout
            logger.debug("Running command: %s", " ".join(cmd))
            return subprocess.run(cmd, capture_output=True, text=True, timeout=to)
        except subprocess.TimeoutExpired as e:
            logger.warning("Command timed out: %s", " ".join(cmd))
            raise
    
    def _ensure_parent_dir(self, path: str):
        Path(path).parent.mkdir(parents=True, exist_ok=True)
    
    # Conversion strategies
    def convert_with_libreoffice(self, docx_path: str, output_path: str) -> bool:
        """Convert DOCX to PDF using LibreOffice in headless mode."""
        if not self._is_tool_available("libreoffice"):
            logger.warning("LibreOffice not available")
            return False
        try:
            self._ensure_parent_dir(output_path)
            cmd = [
                "libreoffice", "--headless", "--convert-to", "pdf",
                "--outdir", str(Path(output_path).parent), docx_path
            ]
            result = self._run_command(cmd)
            if result.returncode == 0:
                expected = Path(output_path).parent / (Path(docx_path).stem + ".pdf")
                if expected.exists():
                    if expected != Path(output_path):
                        shutil.move(str(expected), output_path)
                    logger.info("LibreOffice conversion succeeded")
                    return True
            logger.info("LibreOffice conversion failed: %s", result.stderr)
            return False
        except Exception as e:
            logger.exception("Error with LibreOffice: %s", e)
            return False
    
    def convert_with_unoconv(self, docx_path: str, output_path: str) -> bool:
        """Convert DOCX to PDF using unoconv."""
        if not self._is_tool_available("unoconv"):
            logger.warning("unoconv not available")
            return False
        try:
            self._ensure_parent_dir(output_path)
            cmd = ["unoconv", "-f", "pdf", "-o", output_path, docx_path]
            result = self._run_command(cmd)
            if result.returncode == 0 and Path(output_path).exists():
                logger.info("unoconv conversion succeeded")
                return True
            logger.info("unoconv conversion failed: %s", result.stderr)
            return False
        except Exception as e:
            logger.exception("Error with unoconv: %s", e)
            return False
    
    def convert_with_python_docx_manipulation(self, docx_path: str, output_path: str) -> bool:
        """Modify DOCX structure (margins, fonts, spacing) before conversion."""
        try:
            from docx import Document
            from docx.shared import Inches, Pt
        except ImportError:
            logger.warning("python-docx not available")
            return False
        try:
            doc = Document(docx_path)
            # Reduce margins
            for section in doc.sections:
                section.top_margin = Inches(0.3)
                section.bottom_margin = Inches(0.3)
                section.left_margin = Inches(0.5)
                section.right_margin = Inches(0.5)
            # Reduce font sizes and paragraph spacing
            for paragraph in doc.paragraphs:
                for run in paragraph.runs:
                    if run.font.size:
                        current_size = run.font.size.pt
                        run.font.size = Pt(max(8, current_size * 0.9))
                pf = paragraph.paragraph_format
                pf.space_before = Pt(0)
                pf.space_after = Pt(2)
                pf.line_spacing = 1.0
            # Optimize tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                if run.font.size:
                                    current_size = run.font.size.pt
                                    run.font.size = Pt(max(7, current_size * 0.85))
            with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
                temp_docx = tmp.name
            doc.save(temp_docx)
            success = self.convert_with_libreoffice(temp_docx, output_path)
            try:
                os.remove(temp_docx)
            except Exception:
                pass
            if success:
                logger.info("python-docx manipulation + conversion succeeded")
            return success
        except Exception as e:
            logger.exception("Error with python-docx manipulation: %s", e)
            return False
    
    def get_pdf_page_count(self, pdf_path: str) -> Optional[int]:
        """Return the number of pages in a PDF, or None on error."""
        try:
            import pypdf
            with open(pdf_path, "rb") as f:
                reader = pypdf.PdfReader(f)
                return len(reader.pages)
        except Exception as e:
            logger.warning("Could not determine PDF page count: %s", e)
            return None
    
    def convert_docx_to_pdf_optimized(self, docx_path: str, output_path: str) -> bool:
        """Try multiple strategies and pick the one with fewest pages (ideally 1)."""
        strategies = [
            ("python-docx manipulation", self.convert_with_python_docx_manipulation),
            ("unoconv", self.convert_with_unoconv),
            ("libreoffice", self.convert_with_libreoffice),
        ]
        best_result = None
        best_page_count = float("inf")
        with tempfile.TemporaryDirectory() as tmpdir:
            for name, func in strategies:
                try:
                    safe_name = name.replace(" ", "_")
                    temp_output = os.path.join(tmpdir, f"{safe_name}.pdf")
                    logger.info("Trying strategy: %s", name)
                    success = func(docx_path, temp_output)
                    if not success:
                        logger.info("Strategy '%s' returned False", name)
                        continue
                    if not Path(temp_output).exists():
                        logger.info("Strategy '%s' reported success but file not present", name)
                        continue
                    page_count = self.get_pdf_page_count(temp_output)
                    logger.info("Strategy '%s' produced %s pages", name, page_count)
                    if page_count is None:
                        # treat as failure
                        continue
                    if page_count < best_page_count:
                        # remove previous best
                        if best_result and Path(best_result).exists():
                            try:
                                os.remove(best_result)
                            except Exception:
                                pass
                        # copy current to keep it
                        kept = os.path.join(tmpdir, f"best_{safe_name}.pdf")
                        shutil.copy2(temp_output, kept)
                        best_result = kept
                        best_page_count = page_count
                        if page_count == 1:
                            logger.info("Achieved single-page PDF with strategy: %s", name)
                            break
                    else:
                        logger.info("Strategy '%s' did not improve page count", name)
                except Exception as e:
                    logger.exception("Error while running strategy %s: %s", name, e)
                    continue
            if best_result and Path(best_result).exists():
                shutil.move(best_result, output_path)
                logger.info("Best conversion achieved %s pages", best_page_count)
                return True
        # As a final fallback, try a plain libreoffice conversion directly to output_path
        logger.info("All optimization strategies failed or produced no improvement; trying LibreOffice fallback")
        return self.convert_with_libreoffice(docx_path, output_path)

def convert_docx_to_pdf_advanced(docx_path: str, output_path: str) -> bool:
    """Public function to convert DOCX to PDF with advanced optimization.
    This is the main entry point for the enhanced conversion."""
    converter = AdvancedDocxConverter()
    return converter.convert_docx_to_pdf_optimized(docx_path, output_path)
